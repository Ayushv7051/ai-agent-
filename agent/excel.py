"""
Excel Spreadsheet (.xlsx) Generator for Ayush AI using openpyxl.
Applies automated column sizing, header styling, freeze panes, and zebra fills.
"""

from typing import List, Any, Dict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_DIR
from agent.permissions import validate_safe_path
from agent.logger import logger


def create_excel(filename: str, headers: List[str], rows: List[List[Any]]) -> Dict[str, Any]:
    """
    Creates a styled Microsoft Excel spreadsheet (.xlsx) in files/excel/.
    """
    clean_name = filename.strip()
    if not clean_name.endswith(".xlsx"):
        clean_name += ".xlsx"

    target_path = validate_safe_path(EXCEL_DIR / clean_name, allowed_roots=[EXCEL_DIR])
    logger.info(f"Generating Excel file '{clean_name}' at {target_path}")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Styling Tokens
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10)
        data_align = Alignment(vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # Write Headers
        ws.append(headers)
        ws.row_dimensions[1].height = 26

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Write Rows safely
        row_count = 0
        for r_idx, row_data in enumerate(rows, start=2):
            safe_row = ["" if v is None else str(v) for v in row_data]
            # Ensure row length matches headers
            if len(safe_row) < len(headers):
                safe_row.extend([""] * (len(headers) - len(safe_row)))
            elif len(safe_row) > len(headers):
                safe_row = safe_row[:len(headers)]

            ws.append(safe_row)
            ws.row_dimensions[r_idx].height = 20
            row_count += 1

            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill

        # Freeze Header Row
        ws.freeze_panes = "A2"

        # Auto-fit column widths with padding
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            adjusted_width = max(max_len + 4, 12)
            # Cap maximum width to keep sheet scannable
            ws.column_dimensions[col_letter].width = min(adjusted_width, 45)

        wb.save(target_path)
        logger.info(f"Excel file created successfully with {row_count} rows: {target_path}")

        if target_path.exists() and target_path.stat().st_size > 0:
            size_b = target_path.stat().st_size
            return {
                "success": True,
                "result": {
                    "filename": clean_name,
                    "file_path": str(target_path),
                    "rows_written": row_count,
                    "size_bytes": size_b,
                    "download_url": f"/files/excel/{clean_name}"
                },
                "error": None,
                "verification": {
                    "verified": True,
                    "details": f"Excel spreadsheet verified on disk ({size_b} bytes, {row_count} rows)."
                }
            }
        return {
            "success": False,
            "result": None,
            "error": "Excel spreadsheet was not written to disk.",
            "verification": {"verified": False, "details": "Post-save verification failed."}
        }
    except Exception as e:
        logger.error(f"Failed to create Excel spreadsheet '{clean_name}': {e}")
        return {
            "success": False,
            "result": None,
            "error": f"Excel generation error: {str(e)}",
            "verification": {"verified": False, "details": "Exception during generation."}
        }
